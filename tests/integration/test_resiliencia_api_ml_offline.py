"""Prova de resiliência ponta a ponta — Exercício 24-A (Seção 4.7 / Torneio).

Este é o teste mais importante do exercício: com a API de ML
completamente fora do ar, o fluxo real do webapp (POST
/api/aula22/dashboard, via TestClient sobre webapp.main.app) precisa
continuar respondendo 200, com a 9ª aba "Decisões de ML" mostrando
REVISAO_ML_OFFLINE para quem entrou no ML — nenhum registro Ambíguo
pode ficar de fora, e o processo não pode lançar 500.

A "queda" da API é simulada com monkeypatch em
`ml_client_global.classificar` (o cliente de módulo único usado por
webapp/main.py), não por um MLClient apontando para uma porta fechada
— isso evita depender de timing de rede (timeout de 3s por chamada)
e mantém o teste rápido e determinístico, testando exatamente o mesmo
caminho de fallback que uma porta fechada acionaria (classificar()
devolvendo None).
"""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

import webapp.main as webapp_main

pytestmark = pytest.mark.integration

client = TestClient(webapp_main.app)


def _planilha_com_ambiguos_mapeavel_e_nao_mapeavel() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    insp = wb.create_sheet("Insp_15_06_2026")
    insp.append(["título da planilha — linha de metadados"])
    insp.append(["gerado por: teste"])
    insp.append(["lote_id", "produto", "linha", "turno", "status", "responsavel", "data", "observacao"])
    insp.append(["L001", "TV55", "L1", "A", "OK", "Ana", "15/06/2026", None])
    insp.append(["L002", "TV55", "L1", "A", "EM AJUSTE", "Ana", "15/06/2026", None])  # Ambíguo, mapeável
    insp.append(["L003", "TV55", "L1", "A", "REPROVADO", "Ana", "15/06/2026", "motivo qualquer"])
    insp.append(["L004", "TV55", "L1", "B", "APROVADO PARCIAL", "Ana", "15/06/2026", None])  # Ambíguo, não mapeável
    insp.append(["Total de registros: 4"])

    base = wb.create_sheet("Base_Referencia")
    base.append(["Base de Referência — título"])
    base.append(["lote_id", "codigo_produto", "descricao_produto", "status_cadastro"])
    for lote in ("L001", "L002", "L003", "L004"):
        base.append([lote, "TV55", "Televisão 55pol", "Ativo"])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def upload_com_ambiguos():
    return {
        "arquivo": (
            "inspecao_lotes_10dias.xlsx",
            _planilha_com_ambiguos_mapeavel_e_nao_mapeavel(),
            "application/octet-stream",
        )
    }


@pytest.fixture
def api_ml_totalmente_fora_do_ar(monkeypatch):
    """Simula a API de ML completamente indisponível: toda chamada a
    ml_client_global.classificar(...) devolve None, exatamente como uma
    porta fechada/timeout resultaria depois do tratamento de erro de
    MLClient (ver src/ml_client.py) — sem depender de rede real."""
    monkeypatch.setattr(webapp_main.ml_client_global, "classificar", lambda *args, **kwargs: None)


def test_dashboard_continua_200_com_api_de_ml_totalmente_fora_do_ar(
    upload_com_ambiguos, api_ml_totalmente_fora_do_ar
):
    resposta = client.post("/api/aula22/dashboard", files=upload_com_ambiguos)

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total"] == 4
    assert corpo["por_classificacao"]["Ambíguo"] == 2


def test_excel_gerado_tem_9_abas_com_revisao_ml_offline_e_nenhum_ambiguo_de_fora(
    upload_com_ambiguos, api_ml_totalmente_fora_do_ar
):
    resposta_post = client.post("/api/aula22/dashboard", files=upload_com_ambiguos)
    dashboard_id = resposta_post.json()["id"]

    resposta_download = client.get(f"/api/aula22/dashboard/{dashboard_id}/download")
    assert resposta_download.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(resposta_download.content))
    assert wb.sheetnames == [
        "Resumo", "Todos", "Válidos", "Divergências", "Ambíguos", "Erros de Entrada",
        "Ranking de Regras", "Dicionário", "Decisões de ML",
    ]

    ws_decisoes = wb["Decisões de ML"]
    linhas_decisoes_ml = [
        {
            "lote_id": ws_decisoes.cell(row=r, column=1).value,
            "entrou_no_ml": ws_decisoes.cell(row=r, column=2).value,
            "classe_ml": ws_decisoes.cell(row=r, column=3).value,
            "motivo": ws_decisoes.cell(row=r, column=7).value,
        }
        for r in range(2, ws_decisoes.max_row + 1)
    ]

    # Paridade: nenhum registro Ambíguo ficou de fora, agora no caminho
    # real via API HTTP (não só chamando a função direto, como no Commit 5).
    linhas_ambiguos = wb["Ambíguos"].max_row - 1
    assert len(linhas_decisoes_ml) == linhas_ambiguos == 2

    # L002 (EM AJUSTE) é mapeável: entrou no ML, e a API estando fora do
    # ar tem que virar REVISAO_ML_OFFLINE — nunca uma exceção, nunca None
    # "solto" na aba.
    decisao_l002 = next(item for item in linhas_decisoes_ml if item["lote_id"] == "L002")
    assert decisao_l002["entrou_no_ml"] is True
    assert decisao_l002["classe_ml"] == "REVISAO_ML_OFFLINE"

    # L004 (APROVADO PARCIAL) não é mapeável: nem chega a tentar o ML —
    # isso é "não aplicável", não o fallback de falha (ver src/item_processor.py).
    decisao_l004 = next(item for item in linhas_decisoes_ml if item["lote_id"] == "L004")
    assert decisao_l004["entrou_no_ml"] is False
    assert decisao_l004["classe_ml"] is None
