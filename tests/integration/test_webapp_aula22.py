"""Testes dos endpoints /api/aula22/dashboard (webapp/main.py).

Constrói uma planilha sintética no formato que `carregar_planilha_10dias`
espera (título + metadados antes do cabeçalho nas abas diárias, título
antes do cabeçalho na Base_Referencia) — não depende do dataset real.
"""

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from webapp.main import app

pytestmark = pytest.mark.integration

client = TestClient(app)


def _planilha_sintetica() -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    insp = wb.create_sheet("Insp_15_06_2026")
    insp.append(["título da planilha — linha de metadados"])
    insp.append(["gerado por: teste"])
    insp.append(["lote_id", "produto", "linha", "turno", "status", "responsavel", "data", "observacao"])
    insp.append(["L001", "TV55", "L1", "A", "OK", "Ana", "15/06/2026", None])
    insp.append(["L002", "TV55", "L1", "A", "EM AJUSTE", "Ana", "15/06/2026", None])
    insp.append(["L003", "TV55", "L1", "A", "REPROVADO", "Ana", "15/06/2026", None])
    insp.append(["Total de registros: 3"])

    base = wb.create_sheet("Base_Referencia")
    base.append(["Base de Referência — título"])
    base.append(["lote_id", "codigo_produto", "descricao_produto", "status_cadastro"])
    base.append(["L001", "TV55", "Televisão 55pol", "Ativo"])
    base.append(["L002", "TV55", "Televisão 55pol", "Ativo"])
    base.append(["L003", "TV55", "Televisão 55pol", "Ativo"])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def upload_ok():
    return {"arquivo": ("inspecao_lotes_10dias.xlsx", _planilha_sintetica(), "application/octet-stream")}


def test_post_dashboard_retorna_resumo_com_id(upload_ok):
    resposta = client.post("/api/aula22/dashboard", files=upload_ok)

    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total"] == 3
    assert corpo["por_classificacao"]["Válido"] == 1  # L001 (OK -> APROVADO)
    assert corpo["por_classificacao"]["Ambíguo"] == 1  # L002 (EM AJUSTE)
    assert corpo["por_classificacao"]["Divergência"] == 1  # L003 (REPROVADO sem observação)
    assert corpo["evolucao_por_dia"] == [
        {"dia": "15/06/2026", "total": 3, "Válido": 1, "Divergência": 1, "Ambíguo": 1, "Erro de Entrada": 0}
    ]
    assert "id" in corpo


def test_download_devolve_xlsx_com_9_abas(upload_ok):
    """9 abas desde o Commit 5 (Exercício 24-A): criar_dashboard() sempre
    calcula decisoes_ml (mesmo que vazia) e passa para
    gerar_relatorio_aula22, então a aba "Decisões de ML" está sempre
    presente a partir do webapp — diferente de chamar
    gerar_relatorio_aula22() diretamente sem o parâmetro (8 abas, ver
    tests/integration/test_relatorio_decisoes_ml.py)."""
    resposta_post = client.post("/api/aula22/dashboard", files=upload_ok)
    dashboard_id = resposta_post.json()["id"]

    resposta = client.get(f"/api/aula22/dashboard/{dashboard_id}/download")

    assert resposta.status_code == 200
    assert resposta.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(resposta.content))
    assert wb.sheetnames == [
        "Resumo", "Todos", "Válidos", "Divergências", "Ambíguos", "Erros de Entrada",
        "Ranking de Regras", "Dicionário", "Decisões de ML",
    ]
    # L002 (EM AJUSTE) é o único Ambíguo da planilha sintética — deve ter
    # exatamente 1 linha de dados em "Decisões de ML".
    assert wb["Decisões de ML"].max_row == 2


def test_resumo_executivo_devolve_texto_markdown(upload_ok):
    resposta_post = client.post("/api/aula22/dashboard", files=upload_ok)
    dashboard_id = resposta_post.json()["id"]

    resposta = client.get(f"/api/aula22/dashboard/{dashboard_id}/resumo-executivo")

    assert resposta.status_code == 200
    assert "## Visão Geral" in resposta.text
    assert "## Destaque" in resposta.text


def test_resumo_executivo_com_id_inexistente_retorna_404():
    resposta = client.get("/api/aula22/dashboard/id-que-nao-existe/resumo-executivo")
    assert resposta.status_code == 404


def test_excel_e_resumo_executivo_nascem_do_mesmo_objeto_de_indicadores(upload_ok):
    """Seção 5 do enunciado: total de registros e regra mais acionada devem
    ser idênticos no Excel (aba Resumo) e no resumo_executivo.md, porque os
    dois nascem do mesmo OperationalIndicators calculado uma única vez."""
    resposta_post = client.post("/api/aula22/dashboard", files=upload_ok)
    dashboard_id = resposta_post.json()["id"]

    resposta_download = client.get(f"/api/aula22/dashboard/{dashboard_id}/download")
    resposta_resumo = client.get(f"/api/aula22/dashboard/{dashboard_id}/resumo-executivo")

    wb = openpyxl.load_workbook(io.BytesIO(resposta_download.content))
    total_no_excel = wb["Resumo"]["A5"].value
    ranking_no_excel = wb["Ranking de Regras"]
    regra_mais_acionada_no_excel = ranking_no_excel.cell(row=2, column=2).value  # nome legível, 1ª linha do ranking

    texto_resumo = resposta_resumo.text
    assert str(total_no_excel) in texto_resumo
    assert regra_mais_acionada_no_excel in texto_resumo


def test_log_devolve_texto_da_execucao(upload_ok):
    resposta_post = client.post("/api/aula22/dashboard", files=upload_ok)
    dashboard_id = resposta_post.json()["id"]

    resposta = client.get(f"/api/aula22/dashboard/{dashboard_id}/log")

    assert resposta.status_code == 200
    assert "LOG DE EXECUÇÃO" in resposta.text
    assert "Total de registros processados: 3" in resposta.text


def test_download_com_id_inexistente_retorna_404():
    resposta = client.get("/api/aula22/dashboard/id-que-nao-existe/download")
    assert resposta.status_code == 404


def test_log_com_id_inexistente_retorna_404():
    resposta = client.get("/api/aula22/dashboard/id-que-nao-existe/log")
    assert resposta.status_code == 404


def test_formato_nao_suportado_retorna_400():
    resposta = client.post(
        "/api/aula22/dashboard",
        files={"arquivo": ("dados.csv", b"lote_id,produto\nL001,TV", "text/csv")},
    )
    assert resposta.status_code == 400


def test_planilha_sem_aba_base_referencia_retorna_400():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    insp = wb.create_sheet("Insp_15_06_2026")
    insp.append(["título"])
    insp.append(["metadados"])
    insp.append(["lote_id", "produto", "linha", "turno", "status", "responsavel", "data", "observacao"])
    insp.append(["L001", "TV55", "L1", "A", "OK", "Ana", "15/06/2026", None])
    buffer = io.BytesIO()
    wb.save(buffer)

    resposta = client.post(
        "/api/aula22/dashboard",
        files={"arquivo": ("inspecao.xlsx", buffer.getvalue(), "application/octet-stream")},
    )

    assert resposta.status_code == 400
    assert "Base_Referencia" in resposta.json()["detail"]
