"""Testes de `src/aula22_relatorio.py` — geração do .xlsx de 6 abas + dashboard.

Usa uma lista sintética de RegistroValidado (não depende do dataset real
de 10 dias) para validar a estrutura exigida pelo critério de aceite:
nomes/ordem das abas, nenhuma mistura de classificação, soma das
categorias == total, e presença dos gráficos nativos na aba Resumo.
"""

from datetime import datetime
from unittest.mock import patch

import openpyxl
import pytest

from src.aula22_classificacao import RegistroValidado
from src.aula22_relatorio import NOMES_ABA, gerar_relatorio_aula22

pytestmark = pytest.mark.integration


def _registro(dia, data_referencia, lote_id, classificacao, regra, **overrides):
    base = dict(
        dia=dia,
        data_referencia=data_referencia,
        linha_planilha=4,
        lote_id=lote_id,
        produto="TV55",
        linha="L1",
        turno="A",
        status_original="APROVADO",
        status_normalizado="APROVADO",
        responsavel="Ana",
        data_lote=data_referencia,
        observacao="",
        ocorrencia_no_dia=1,
        classificacao=classificacao,
        regra=regra,
        motivo="motivo de teste",
    )
    base.update(overrides)
    return RegistroValidado(**base)


@pytest.fixture
def registros():
    return [
        _registro("Insp_15_06_2026", "15/06/2026", "L001", "Válido", "RN08"),
        _registro("Insp_15_06_2026", "15/06/2026", "L002", "Válido", "RN08"),
        _registro("Insp_15_06_2026", "15/06/2026", "L003", "Divergência", "RN05"),
        _registro("Insp_15_06_2026", "15/06/2026", "L004", "Ambíguo", "RN09"),
        _registro("Insp_16_06_2026", "16/06/2026", "", "Erro de Entrada", "RN01-RN04"),
        _registro("Insp_16_06_2026", "16/06/2026", "L006", "Válido", "RN08"),
    ]


def test_gera_6_abas_com_nomes_e_ordem_corretos(tmp_path, registros):
    caminho = tmp_path / "relatorio_conferencia_lotes.xlsx"
    gerar_relatorio_aula22(registros, str(caminho))

    wb = openpyxl.load_workbook(caminho)
    assert wb.sheetnames == ["Resumo", "Todos", "Válidos", "Divergências", "Ambíguos", "Erros de Entrada"]


@pytest.mark.regression
def test_nenhuma_aba_mistura_classificacoes(tmp_path, registros):
    caminho = tmp_path / "relatorio.xlsx"
    gerar_relatorio_aula22(registros, str(caminho))
    wb = openpyxl.load_workbook(caminho)

    for classificacao, nome_aba in NOMES_ABA.items():
        if classificacao in ("Resumo", "Todos"):
            continue
        ws = wb[nome_aba]
        cabecalho = [c.value for c in ws[1]]
        coluna = cabecalho.index("classificacao") + 1
        valores = {ws.cell(row=r, column=coluna).value for r in range(2, ws.max_row + 1)}
        assert valores == {classificacao}, f"aba {nome_aba} misturou classificações: {valores}"


def test_soma_das_categorias_igual_ao_total(tmp_path, registros):
    caminho = tmp_path / "relatorio.xlsx"
    gerar_relatorio_aula22(registros, str(caminho))
    wb = openpyxl.load_workbook(caminho)

    total_categorias = sum(
        wb[nome_aba].max_row - 1
        for classificacao, nome_aba in NOMES_ABA.items()
        if classificacao not in ("Resumo", "Todos")
    )
    assert total_categorias == wb["Todos"].max_row - 1 == len(registros)


def test_aba_resumo_tem_graficos_nativos(tmp_path, registros):
    caminho = tmp_path / "relatorio.xlsx"
    gerar_relatorio_aula22(registros, str(caminho))
    wb = openpyxl.load_workbook(caminho)

    tipos = [type(g).__name__ for g in wb["Resumo"]._charts]
    assert "DoughnutChart" in tipos
    assert "LineChart" in tipos


def test_resumo_em_memoria_bate_com_as_abas(tmp_path, registros):
    caminho = tmp_path / "relatorio.xlsx"
    resultado = gerar_relatorio_aula22(registros, str(caminho))

    assert resultado["resumo"]["total"] == len(registros)
    assert resultado["resumo"]["por_classificacao"]["Válido"] == 3
    assert resultado["resumo"]["por_classificacao"]["Divergência"] == 1
    assert resultado["resumo"]["por_classificacao"]["Ambíguo"] == 1
    assert resultado["resumo"]["por_classificacao"]["Erro de Entrada"] == 1
    assert "log" in resultado and "LOG DE EXECUÇÃO" in resultado["log"]


def test_gerar_log_execucao_lista_todos_os_dias(registros):
    from src.aula22_relatorio import gerar_log_execucao

    log = gerar_log_execucao(registros)
    assert "Insp_15_06_2026" in log
    assert "Insp_16_06_2026" in log
    assert "Total de registros processados: 6" in log


def test_log_de_execucao_usa_data_hora_mockada(registros):
    """Mock de datetime.now() garante que o carimbo do log é determinístico.

    Sem mock, o teste só poderia verificar 'contém alguma data' — com mock,
    verificamos o valor exato. Isso impede que o teste fique frágil por
    causa do relógio do sistema (seção 5.5 do enunciado da Aula 23).
    """
    data_fixa = datetime(2026, 6, 30, 8, 0, 0)

    with patch("src.aula22_relatorio.datetime") as mock_dt:
        mock_dt.now.return_value = data_fixa
        mock_dt.strptime = datetime.strptime  # preserva strptime real (usado internamente)

        from src.aula22_relatorio import gerar_log_execucao
        log = gerar_log_execucao(registros)

    assert "30/06/2026 08:00:00" in log
    assert "Total de registros processados: 6" in log


def test_classificacao_com_base_referencia_mockada_via_patch():
    """Demonstra uso de unittest.mock.patch para simular a Base_Referencia
    como se fosse um sistema externo (seção 5.5 do enunciado da Aula 23).

    Em vez de passar um DataFrame real, mockamos as duas funções que
    consultam a base (verificar_existencia_lote e verificar_status_lote)
    para retornar respostas controladas — provando que o pipeline de
    classificação funciona sem depender de arquivo em disco.
    """
    from src.aula22_classificacao import classificar_registro

    registro = {
        "lote_id": "LOTE-MOCK-001",
        "produto": "TV55",
        "linha": "L1",
        "turno": "A",
        "status": "APROVADO",
        "responsavel": "Ana",
        "data": "15/06/2026",
        "observacao": "",
        "_dia": "Insp_15_06_2026",
        "_data_referencia": "15/06/2026",
        "_linha_planilha": 4,
    }

    with patch("src.aula22_classificacao.verificar_existencia_lote", return_value=True) as mock_existe, \
         patch("src.aula22_classificacao.verificar_status_lote", return_value=True) as mock_status:

        resultado = classificar_registro(registro, ocorrencia_no_dia=1, base_referencia=None)

    # O registro é válido porque os mocks dizem que o lote existe e está ativo
    assert resultado.classificacao == "Válido"
    assert resultado.regra == "RN08"

    # Confirma que os mocks foram chamados (a base real nunca foi consultada)
    mock_existe.assert_called_once_with(None, "LOTE-MOCK-001")
    mock_status.assert_called_once_with(None, "LOTE-MOCK-001")
