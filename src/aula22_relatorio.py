"""Gerador do relatorio_conferencia_lotes.xlsx — Aula 22.

Substitui (não estende) o gerar_relatorio() de src/relatorio.py para este
fluxo: aquele escreve via DataFrame.to_excel() direto no ExcelWriter, sem
acesso ao Workbook — não há gancho para inserir gráficos nativos. Aqui
abrimos o Workbook do openpyxl diretamente para poder usar
openpyxl.chart.DoughnutChart e LineChart.

Produz exatamente 8 abas — as 6 tabelas por categoria (nenhuma mistura
Divergência/Ambíguo, conforme o critério de aceite) mais 2 abas de
apoio derivadas dos indicadores operacionais (Aula 24):

    Resumo | Todos | Válidos | Divergências | Ambíguos | Erros de Entrada
    | Ranking de Regras | Dicionário

E devolve também o texto do log de execução (RN de evidência da Aula 1),
que o chamador decide se grava como .txt, retorna na API, ou ambos.
"""
from __future__ import annotations

import io
import logging
from collections import Counter, OrderedDict
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (
    CharacterProperties,
    Paragraph,
    ParagraphProperties,
    RichTextProperties,
)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.aula22_classificacao import RegistroValidado
from src.operational_indicators import (
    REGRAS_DESCRICAO,
    OperationalIndicators,
    calcular_indicadores,
)

logger = logging.getLogger("aula22.relatorio")

NOMES_ABA = {
    "Resumo": "Resumo",
    "Todos": "Todos",
    "Válido": "Válidos",
    "Divergência": "Divergências",
    "Ambíguo": "Ambíguos",
    "Erro de Entrada": "Erros de Entrada",
}

COLUNAS_RELATORIO = [
    "data_referencia", "lote_id", "produto", "linha", "turno",
    "status_original", "status_normalizado", "responsavel",
    "observacao", "classificacao", "regra", "motivo", "linha_planilha",
]

COR_CABECALHO = "1F2937"
COR_TEXTO_CABECALHO = "FFFFFF"
CORES_CLASSIFICACAO = {
    "Válido": "16A34A",
    "Divergência": "DC2626",
    "Ambíguo": "D97706",
    "Erro de Entrada": "6B7280",
}


def _dataframe_por_classificacao(registros: list[RegistroValidado], classificacao: str | None) -> pd.DataFrame:
    if classificacao is None:
        selecionados = registros
    else:
        selecionados = [r for r in registros if r.classificacao == classificacao]
    linhas = [r.to_dict() for r in selecionados]
    df = pd.DataFrame(linhas, columns=[
        "data_referencia", "lote_id", "produto", "linha", "turno",
        "status_original", "status_normalizado", "responsavel",
        "observacao", "classificacao", "regra", "motivo", "linha_planilha",
    ])
    return df


def _escrever_aba_tabela(wb: Workbook, nome_aba: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(nome_aba)
    ws.append(list(df.columns))
    for celula in ws[1]:
        celula.font = Font(bold=True, color=COR_TEXTO_CABECALHO)
        celula.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        celula.alignment = Alignment(horizontal="center")

    for _, linha in df.iterrows():
        ws.append(list(linha))

    for indice, coluna in enumerate(df.columns, start=1):
        maior = max([len(str(coluna))] + [len(str(v)) for v in df[coluna].astype(str)]) if len(df) else len(str(coluna))
        ws.column_dimensions[get_column_letter(indice)].width = min(max(maior + 2, 10), 42)

    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"


def _serie_evolucao_por_dia(registros: list[RegistroValidado]) -> "OrderedDict[str, dict]":
    """Agrupa por _data_referencia (ordem cronológica) e conta por classificação."""
    dias: "OrderedDict[str, dict]" = OrderedDict()
    for r in sorted(registros, key=lambda x: datetime.strptime(x.data_referencia, "%d/%m/%Y")):
        chave = r.data_referencia
        if chave not in dias:
            dias[chave] = {"total": 0, "Válido": 0, "Divergência": 0, "Ambíguo": 0, "Erro de Entrada": 0}
        dias[chave]["total"] += 1
        dias[chave][r.classificacao] += 1
    return dias


def _rotulo_curto_dia(data_referencia: str) -> str:
    """'15/06/2026' -> '15/06'. Datas completas lado a lado no eixo do
    gráfico de evolução se sobrepõem (10 rótulos longos) e viram uma
    linha sólida ilegível — o ano é redundante aqui, já que os 10 dias
    do relatório são todos do mesmo ano."""
    return datetime.strptime(data_referencia, "%d/%m/%Y").strftime("%d/%m")


def _rotacionar_rotulos_eixo(eixo, graus: int = -45) -> None:
    """Inclina os rótulos do eixo em `graus` graus para caberem sem
    sobrepor, quando há muitas categorias (ex.: 10 dias).

    Também reduz o tamanho da fonte dos rótulos (sz=900 = 9pt, contra
    o padrão de ~10-18pt) — sem isso, o texto rotacionado em -45° não
    cabe na área reservada abaixo do plot area e o Excel/LibreOffice
    corta visualmente o final de cada rótulo (ex.: "15/..." em vez de
    "15/06"). Fonte menor = a "caixa" necessária pro texto rotacionado
    fica menor e passa a caber no espaço disponível.
    """
    rot = int(graus * 60000)  # openpyxl espera 1/60000 de grau
    propriedades_corpo = RichTextProperties(rot=rot, vert="horz")
    fonte_rotulo = CharacterProperties(sz=900)
    paragrafo = Paragraph(
        pPr=ParagraphProperties(defRPr=fonte_rotulo),
        endParaRPr=fonte_rotulo,
    )
    eixo.txPr = RichText(bodyPr=propriedades_corpo, p=[paragrafo])


def _escrever_aba_resumo(wb: Workbook, registros: list[RegistroValidado], indicadores: OperationalIndicators) -> None:
    """Os números desta aba vêm de `indicadores` (já calculados por
    calcular_indicadores), nunca recontados aqui — a mesma regra que
    protege o resumo_executivo.md contra divergir do Excel (ver módulo
    src/operational_indicators.py)."""
    ws = wb.create_sheet("Resumo", 0)  # sempre a primeira aba

    contagem_por_classe = {
        "Válido": indicadores.qtd_validos,
        "Divergência": indicadores.qtd_divergencias,
        "Ambíguo": indicadores.qtd_ambiguos,
        "Erro de Entrada": indicadores.qtd_erros_entrada,
    }

    ws["A1"] = "Relatório de Conferência de Lotes — Resumo Executivo"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws["A2"].font = Font(italic=True, color="6B7280")

    # Indicadores numéricos (linha 4 = cabeçalho, linha 5 = valores)
    cabecalhos = ["Total de registros", "Válidos", "% Válidos", "Divergências", "% Divergências",
                  "Ambíguos", "% Ambíguos", "Erros de Entrada", "% Erros de Entrada"]
    valores = [
        indicadores.total_registros,
        indicadores.qtd_validos, round(indicadores.pct_validos, 1),
        indicadores.qtd_divergencias, round(indicadores.pct_divergencias, 1),
        indicadores.qtd_ambiguos, round(indicadores.pct_ambiguos, 1),
        indicadores.qtd_erros_entrada, round(indicadores.pct_erros_entrada, 1),
    ]
    linha_cabecalho = 4
    for indice, (cab, val) in enumerate(zip(cabecalhos, valores), start=1):
        celula_cab = ws.cell(row=linha_cabecalho, column=indice, value=cab)
        celula_cab.font = Font(bold=True, color=COR_TEXTO_CABECALHO)
        celula_cab.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        celula_cab.alignment = Alignment(horizontal="center", wrap_text=True)
        celula_val = ws.cell(row=linha_cabecalho + 1, column=indice, value=val)
        celula_val.alignment = Alignment(horizontal="center")
        celula_val.font = Font(bold=True, size=12)
        ws.column_dimensions[get_column_letter(indice)].width = 16

    # --- Tabela de apoio para o gráfico de rosca (classificação x total) ---
    linha_base_rosca = 8
    ws.cell(row=linha_base_rosca, column=1, value="Classificação").font = Font(bold=True)
    ws.cell(row=linha_base_rosca, column=2, value="Total").font = Font(bold=True)
    classes_ordem = ["Válido", "Divergência", "Ambíguo", "Erro de Entrada"]
    for offset, classe in enumerate(classes_ordem, start=1):
        ws.cell(row=linha_base_rosca + offset, column=1, value=classe)
        ws.cell(row=linha_base_rosca + offset, column=2, value=contagem_por_classe[classe])

    rosca = DoughnutChart()
    rosca.title = "Distribuição por classificação"
    dados_rosca = Reference(ws, min_col=2, min_row=linha_base_rosca, max_row=linha_base_rosca + len(classes_ordem))
    categorias_rosca = Reference(ws, min_col=1, min_row=linha_base_rosca + 1, max_row=linha_base_rosca + len(classes_ordem))
    rosca.add_data(dados_rosca, titles_from_data=True)
    rosca.set_categories(categorias_rosca)
    rosca.height = 8
    rosca.width = 12
    # Só percentual no rótulo — sem isso, o openpyxl deixa
    # showCatName/showSerName/showVal indefinidos e o Excel/LibreOffice
    # os preenche como True, resultando em rótulos poluídos do tipo
    # "Total; Válido; 150; 60%" em vez de só "60%".
    rosca.dataLabels = DataLabelList(
        showPercent=True,
        showCatName=False,
        showSerName=False,
        showVal=False,
        showLegendKey=False,
        showBubbleSize=False,
    )
    # Obs.: rosca só é ancorada na planilha mais abaixo, junto com o
    # gráfico de evolução, depois que a tabela "Dia" for escrita (ver
    # linha_ancora_graficos) — assim garantimos que nenhum gráfico
    # sobrepõe as tabelas de apoio.

    # --- Tabela de apoio para o gráfico de evolução por dia ---
    serie_dias = _serie_evolucao_por_dia(registros)
    linha_base_evolucao = linha_base_rosca + len(classes_ordem) + 2  # espaço após a tabela da rosca
    cabecalhos_evolucao = ["Dia", "Total", "Válido", "Divergência", "Ambíguo", "Erro de Entrada", "Divergência + Ambíguo"]
    for indice, cab in enumerate(cabecalhos_evolucao, start=1):
        ws.cell(row=linha_base_evolucao, column=indice, value=cab).font = Font(bold=True)

    for offset, (dia, valores_dia) in enumerate(serie_dias.items(), start=1):
        linha = linha_base_evolucao + offset
        ws.cell(row=linha, column=1, value=_rotulo_curto_dia(dia))
        ws.cell(row=linha, column=2, value=valores_dia["total"])
        ws.cell(row=linha, column=3, value=valores_dia["Válido"])
        ws.cell(row=linha, column=4, value=valores_dia["Divergência"])
        ws.cell(row=linha, column=5, value=valores_dia["Ambíguo"])
        ws.cell(row=linha, column=6, value=valores_dia["Erro de Entrada"])
        ws.cell(row=linha, column=7, value=valores_dia["Divergência"] + valores_dia["Ambíguo"])

    ultima_linha_evolucao = linha_base_evolucao + len(serie_dias)

    linha_grafico = LineChart()
    linha_grafico.title = "Evolução dos registros por dia"
    linha_grafico.style = 10
    linha_grafico.y_axis.title = "Registros"
    linha_grafico.y_axis.scaling.min = 0
    linha_grafico.x_axis.title = "Dia"
    linha_grafico.x_axis.delete = False
    # Série obrigatória: Divergência + Ambíguo (coluna 7) — revela se o problema piora/melhora
    dados_evolucao = Reference(
        ws, min_col=7, max_col=7, min_row=linha_base_evolucao, max_row=ultima_linha_evolucao
    )
    # Série extra de contexto: total de registros no dia (coluna 2)
    dados_total = Reference(
        ws, min_col=2, max_col=2, min_row=linha_base_evolucao, max_row=ultima_linha_evolucao
    )
    categorias_evolucao = Reference(ws, min_col=1, min_row=linha_base_evolucao + 1, max_row=ultima_linha_evolucao)
    linha_grafico.add_data(dados_total, titles_from_data=True)
    linha_grafico.add_data(dados_evolucao, titles_from_data=True)
    linha_grafico.set_categories(categorias_evolucao)
    # 10 rótulos de dia lado a lado se sobrepõem se ficarem na
    # horizontal — inclinar evita que virem uma linha sólida ilegível.
    _rotacionar_rotulos_eixo(linha_grafico.x_axis)
    linha_grafico.height = 10
    linha_grafico.width = 18

    # Âncora dos gráficos: sempre abaixo da última linha escrita (tabela
    # "Dia"), calculada dinamicamente. Antes estava fixa em "A14"/"D14",
    # que é justamente onde a tabela "Dia" é escrita (linha_base_evolucao
    # também é 14) — por isso os dois gráficos apareciam flutuando em
    # cima da tabela e um em cima do outro (A e D ficam muito próximas
    # para a largura de 12-18cm de cada gráfico).
    linha_ancora_graficos = ultima_linha_evolucao + 3
    ws.add_chart(rosca, f"A{linha_ancora_graficos}")
    ws.add_chart(linha_grafico, f"H{linha_ancora_graficos}")

    ws.column_dimensions["A"].width = 20


def _escrever_aba_ranking_regras(wb: Workbook, indicadores: OperationalIndicators) -> None:
    """Aba de apoio ao indicador 6 (regra mais acionada): todas as regras
    de divergência/ambiguidade/erro que apareceram no período, da mais
    para a menos frequente. RN08 (Válido) não entra aqui pelo mesmo motivo
    documentado em calcular_indicadores() — ela não é um problema a apontar.
    """
    ws = wb.create_sheet("Ranking de Regras")
    ws.append(["Código", "Regra", "Quantidade", "% do total"])
    for celula in ws[1]:
        celula.font = Font(bold=True, color=COR_TEXTO_CABECALHO)
        celula.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        celula.alignment = Alignment(horizontal="center")

    total = indicadores.total_registros
    for codigo, quantidade in indicadores.ranking_regras:
        nome = REGRAS_DESCRICAO.get(codigo, codigo)
        percentual = round(100 * quantidade / total, 1) if total else 0
        ws.append([codigo, nome, quantidade, percentual])

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.freeze_panes = "A2"
    if indicadores.ranking_regras:
        ws.auto_filter.ref = f"A1:D{len(indicadores.ranking_regras) + 1}"


def _escrever_aba_dicionario(wb: Workbook) -> None:
    """Aba de referência fixa: o significado de cada código de regra (RN)
    que pode aparecer na coluna "regra" das demais abas — para quem lê o
    relatório sem ter o enunciado das RNs em mãos.
    """
    ws = wb.create_sheet("Dicionário")
    ws.append(["Código", "Descrição"])
    for celula in ws[1]:
        celula.font = Font(bold=True, color=COR_TEXTO_CABECALHO)
        celula.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        celula.alignment = Alignment(horizontal="center")

    for codigo, descricao in REGRAS_DESCRICAO.items():
        ws.append([codigo, descricao])

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 70
    ws.freeze_panes = "A2"


def gerar_relatorio_aula22(
    registros: list[RegistroValidado],
    caminho_saida: str,
    indicadores: OperationalIndicators | None = None,
) -> dict:
    """Gera o .xlsx de 8 abas + dashboard nativo e devolve um resumo em memória.

    Args:
        registros: lista de RegistroValidado já classificados.
        caminho_saida: caminho do .xlsx a ser gravado.
        indicadores: OperationalIndicators já calculado a partir dos mesmos
            `registros` (ver src/operational_indicators.py). Se não for
            informado, é calculado aqui — mas o padrão esperado pelo
            chamador (webapp/main.py) é calcular uma única vez e passar
            adiante, para o Excel e o resumo_executivo.md nascerem do
            mesmo objeto e nunca divergirem entre si.

    Returns:
        dict com "resumo" (contagens/percentuais), "arquivo" (caminho) e
        "log" (texto do log de execução, seção 5.4 do enunciado).
    """
    if indicadores is None:
        indicadores = calcular_indicadores(registros)

    wb = Workbook()
    wb.remove(wb.active)  # remove a aba default "Sheet"

    _escrever_aba_resumo(wb, registros, indicadores)
    _escrever_aba_tabela(wb, "Todos", _dataframe_por_classificacao(registros, None))
    for classificacao, nome_aba in NOMES_ABA.items():
        if classificacao in ("Resumo", "Todos"):
            continue
        _escrever_aba_tabela(wb, nome_aba, _dataframe_por_classificacao(registros, classificacao))
    _escrever_aba_ranking_regras(wb, indicadores)
    _escrever_aba_dicionario(wb)

    wb.save(caminho_saida)
    logger.info("Relatório salvo em %s", caminho_saida)

    resumo = {
        "total": indicadores.total_registros,
        "por_classificacao": {
            "Válido": indicadores.qtd_validos,
            "Divergência": indicadores.qtd_divergencias,
            "Ambíguo": indicadores.qtd_ambiguos,
            "Erro de Entrada": indicadores.qtd_erros_entrada,
        },
        "percentual": {
            "Válido": round(indicadores.pct_validos, 1),
            "Divergência": round(indicadores.pct_divergencias, 1),
            "Ambíguo": round(indicadores.pct_ambiguos, 1),
            "Erro de Entrada": round(indicadores.pct_erros_entrada, 1),
        },
        "evolucao_por_dia": _serie_evolucao_por_dia(registros),
    }

    texto_log = gerar_log_execucao(registros)
    return {"resumo": resumo, "arquivo": caminho_saida, "log": texto_log}


def gerar_log_execucao(registros: list[RegistroValidado]) -> str:
    """Log simples de evidência (seção 5.4): data/hora, total e totais por classificação."""
    contagem = Counter(r.classificacao for r in registros)
    linhas = [
        "=== LOG DE EXECUÇÃO — Conferência de Lotes (Aula 22) ===",
        f"Data/hora da execução: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}",
        f"Total de registros processados: {len(registros)}",
        "",
        "Totais por classificação:",
    ]
    for classe in ("Válido", "Divergência", "Ambíguo", "Erro de Entrada"):
        linhas.append(f"  {classe}: {contagem[classe]}")

    dias = sorted({r.dia for r in registros})
    linhas.append("")
    linhas.append(f"Abas processadas ({len(dias)}): {', '.join(dias)}")
    return "\n".join(linhas)